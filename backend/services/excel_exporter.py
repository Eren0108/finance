import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.config import EXPORT_DIR
from backend.services.analytics import get_monthly_summary
from backend.database import get_db
from backend.utils.helpers import get_month_date_range, format_month_label


def generate_excel_report(month_str: str) -> str:
    """
    Generates a professional Excel (.xlsx) workbook for the specified month
    with 'Summary' and 'Transactions' sheets, styled according to financial reporting standards.
    Returns the absolute path to the generated file.
    """
    summary = get_monthly_summary(month_str)
    start_date, end_date = get_month_date_range(month_str)

    # Fetch all transactions for this month
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT 
                t.date, t.description, t.amount, t.type, t.source,
                c.name as category_name
            FROM transactions t
            JOIN categories c ON t.category_id = c.id
            WHERE t.date >= ? AND t.date <= ?
            ORDER BY t.date ASC, t.id ASC;
        """, (start_date, end_date))
        transactions = cursor.fetchall()

    wb = openpyxl.Workbook()
    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_title = Font(name="Segoe UI", size=16, bold=True, color="1E293B")
    font_subtitle = Font(name="Segoe UI", size=11, italic=True, color="64748B")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
    font_regular = Font(name="Segoe UI", size=11, color="1E293B")
    font_income = Font(name="Segoe UI", size=11, bold=True, color="16A34A")
    font_expense = Font(name="Segoe UI", size=11, bold=True, color="DC2626")

    fill_primary = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    fill_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_highlight = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    double_bottom_border = Border(
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="double", color="1E293B")
    )

    # 1. Report Header
    month_display = format_month_label(month_str)
    ws_summary["A1"] = f"Personal Finance Report - {month_display}"
    ws_summary["A1"].font = font_title
    ws_summary["A2"] = f"Period: {start_date} to {end_date} | Currency: INR (₹)"
    ws_summary["A2"].font = font_subtitle

    # 2. Key Metrics Table
    ws_summary["A4"] = "Monthly Overview"
    ws_summary["A4"].font = font_bold

    metrics = [
        ("Total Income", summary["total_income"], "₹#,##0.00", font_income),
        ("Total Expenses", summary["total_expense"], "₹#,##0.00", font_expense),
        ("Net Savings", summary["savings"], "₹#,##0.00", font_income if summary["savings"] >= 0 else font_expense),
        ("Savings Rate", summary["savings_rate"] / 100.0, "0.0%", font_bold),
        ("Total Transactions", summary["transaction_count"], "#,##0", font_regular)
    ]

    ws_summary.append([]) # row 5
    for idx, (label, val, num_fmt, f_style) in enumerate(metrics, start=5):
        cell_lbl = ws_summary.cell(row=idx, column=1, value=label)
        cell_val = ws_summary.cell(row=idx, column=2, value=val)
        
        cell_lbl.font = font_regular
        cell_lbl.border = thin_border
        cell_lbl.fill = fill_highlight

        cell_val.font = f_style
        cell_val.number_format = num_fmt
        cell_val.border = thin_border
        cell_val.alignment = Alignment(horizontal="right")

    # 3. Category Breakdown Table
    cat_start_row = 12
    ws_summary.cell(row=cat_start_row, column=1, value="Category-wise Spending Breakdown").font = font_bold

    headers_cat = ["Category", "Budget Limit (₹)", "Actual Spent (₹)", "% of Budget", "% of Total Expenses"]
    for col_idx, h in enumerate(headers_cat, start=1):
        c = ws_summary.cell(row=cat_start_row + 1, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx > 3 else "left")
        c.border = thin_border

    curr_row = cat_start_row + 2
    for cat in summary["categories"]:
        c1 = ws_summary.cell(row=curr_row, column=1, value=cat["category_name"])
        c2 = ws_summary.cell(row=curr_row, column=2, value=cat["budget_limit"])
        c3 = ws_summary.cell(row=curr_row, column=3, value=cat["spent"])
        c4 = ws_summary.cell(row=curr_row, column=4, value=(cat["budget_percentage"] / 100.0) if cat["budget_limit"] > 0 else 0)
        c5 = ws_summary.cell(row=curr_row, column=5, value=cat["percentage_of_total"] / 100.0)

        c1.font = font_regular
        c2.font = font_regular
        c3.font = font_bold
        c4.font = font_regular
        c5.font = font_regular

        c2.number_format = "₹#,##0.00"
        c3.number_format = "₹#,##0.00"
        c4.number_format = "0.0%"
        c5.number_format = "0.0%"

        c2.alignment = Alignment(horizontal="right")
        c3.alignment = Alignment(horizontal="right")
        c4.alignment = Alignment(horizontal="right")
        c5.alignment = Alignment(horizontal="right")

        for c in (c1, c2, c3, c4, c5):
            c.border = thin_border
            if curr_row % 2 == 0:
                c.fill = fill_zebra

        curr_row += 1

    # Total row for categories
    ws_summary.cell(row=curr_row, column=1, value="Total").font = font_bold
    ws_summary.cell(row=curr_row, column=1).border = double_bottom_border
    
    tot_limit_cell = ws_summary.cell(row=curr_row, column=2, value=f"=SUM(B{cat_start_row + 2}:B{curr_row - 1})")
    tot_limit_cell.font = font_bold
    tot_limit_cell.number_format = "₹#,##0.00"
    tot_limit_cell.border = double_bottom_border

    tot_spent_cell = ws_summary.cell(row=curr_row, column=3, value=f"=SUM(C{cat_start_row + 2}:C{curr_row - 1})")
    tot_spent_cell.font = font_bold
    tot_spent_cell.number_format = "₹#,##0.00"
    tot_spent_cell.border = double_bottom_border

    ws_summary.cell(row=curr_row, column=4, value="").border = double_bottom_border
    ws_summary.cell(row=curr_row, column=5, value=1.0).font = font_bold
    ws_summary.cell(row=curr_row, column=5).number_format = "0.0%"
    ws_summary.cell(row=curr_row, column=5).border = double_bottom_border

    # Sheet 2: Transactions
    ws_tx = wb.create_sheet(title="Transactions")
    ws_tx.views.sheetView[0].showGridLines = True

    tx_headers = ["Date", "Description", "Category", "Type", "Amount (₹)", "Source"]
    for col_idx, h in enumerate(tx_headers, start=1):
        c = ws_tx.cell(row=1, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_primary
        c.alignment = Alignment(horizontal="center" if h in ("Date", "Type", "Source") else ("right" if "Amount" in h else "left"))
        c.border = thin_border

    row_num = 2
    for t in transactions:
        c_date = ws_tx.cell(row=row_num, column=1, value=t["date"])
        c_desc = ws_tx.cell(row=row_num, column=2, value=t["description"])
        c_cat = ws_tx.cell(row=row_num, column=3, value=t["category_name"])
        c_type = ws_tx.cell(row=row_num, column=4, value=t["type"].capitalize())
        c_amt = ws_tx.cell(row=row_num, column=5, value=float(t["amount"]))
        c_src = ws_tx.cell(row=row_num, column=6, value=t["source"].capitalize())

        c_date.font = font_regular
        c_desc.font = font_regular
        c_cat.font = font_regular
        c_type.font = font_income if t["type"] == "income" else font_expense
        c_amt.font = font_income if t["type"] == "income" else font_expense
        c_src.font = font_regular

        c_date.alignment = Alignment(horizontal="center")
        c_type.alignment = Alignment(horizontal="center")
        c_amt.alignment = Alignment(horizontal="right")
        c_amt.number_format = "₹#,##0.00"
        c_src.alignment = Alignment(horizontal="center")

        for c in (c_date, c_desc, c_cat, c_type, c_amt, c_src):
            c.border = thin_border
            if row_num % 2 == 1:
                c.fill = fill_zebra

        row_num += 1

    # Auto-fit column widths for both sheets
    for ws in (ws_summary, ws_tx):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if val.startswith("="):
                    val = "123,456.78"
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Save workbook
    output_filename = f"finance_report_{month_str}.xlsx"
    output_path = EXPORT_DIR / output_filename
    wb.save(output_path)
    return str(output_path)
